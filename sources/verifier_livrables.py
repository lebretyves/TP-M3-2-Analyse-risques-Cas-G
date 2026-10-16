"""Vérifie les livrables sans modifier le classeur ; openpyxl est utilisé en lecture."""
from pathlib import Path
import hashlib, json
from collections import Counter
import openpyxl, pymupdf
ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'sources/risques.json').read_text(encoding='utf-8'))
src=openpyxl.load_workbook(ROOT/'supports/TP_M3_2_Analyse_risques.xlsx')
dst=openpyxl.load_workbook(ROOT/'analyses/m3/analyse_risques_eq02.xlsx')
val=openpyxl.load_workbook(ROOT/'analyses/m3/analyse_risques_eq02.xlsx',data_only=True)
count=0
for name in src.sheetnames:
 for row in src[name]:
  for c in row:
   if c.data_type=='f':
    assert dst[name][c.coordinate].value==c.value,(name,c.coordinate,'formule modifiée')
    count+=1
for i,r in enumerate(data['risques'],9):
 sh=val['Analyse de risques']
 assert all(sh.cell(i,c).value is not None for c in range(1,19))
 assert sh.cell(i,9).value==r['g']*r['p']
 assert sh.cell(i,16).value==r['ga']*r['pa']
 assert sh.cell(i,18).value==r['justification']
checks=[val['Synthèse'].cell(i,2).value for i in range(12,17)]
assert checks==[10,5,3,2,0],checks
for sheet in val:
 for row in sheet:
  for c in row:assert c.data_type!='e',(sheet.title,c.coordinate,c.value)
assert dst['Analyse de risques'].freeze_panes=='D9'
assert len(dst['Analyse de risques'].data_validations.dataValidation)==3
pdf=pymupdf.open(ROOT/'analyses/m3/TP_M3_2_analyse_risques_cas_G.pdf')
assert len(pdf)==5
links=[a for page in pdf for a in page.get_links()]
assert len(links)==9
def level(x):return 'élevé' if x>=6 else 'moyen' if x>=3 else 'faible'
report={'date':'2026-10-16','equipe':2,'risques':10,'formules_originales_preservees':count,'controles_synthese':checks,'avant':dict(Counter(level(r['g']*r['p']) for r in data['risques'])),'residuels_cibles':dict(Counter(level(r['ga']*r['pa']) for r in data['risques'])),'pages_pdf':len(pdf),'liens_pdf':len(links),'essais_materiels_realises':False,'supports_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'supports').iterdir() if p.is_file()}}
(ROOT/'analyses/m3/verification_livrables.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
